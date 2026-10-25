"""Render a document (structured JSON) to .docx with python-docx.

Follows the format spec's section order and labels; Vietnamese-safe (all text is
Unicode; Times New Roman is set explicitly because it carries full diacritics).
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .formats import resolve_title

_FONT = "Times New Roman"


def render_docx(document: dict, format_spec: dict, out_path: Path) -> Path:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = _FONT
    style.font.size = Pt(12)

    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(resolve_title(format_spec, document))
    run.bold = True
    run.font.size = Pt(14)

    if document.get("lesson"):
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.add_run(f"Bài dạy: {document['lesson']}").bold = True
    if document.get("week"):
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.add_run(f"Tuần: {document['week']}")

    sections = document.get("sections", {})
    for spec_section in format_spec.get("sections", []):
        sid = spec_section["id"]
        content = sections.get(sid)
        if content in (None, "", [], {}):
            continue
        _render_section(doc, spec_section, content)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def _render_section(doc: Document, spec_section: dict, content) -> None:
    stype = spec_section.get("type", "text")
    label = spec_section.get("label", "")
    if label:
        heading = doc.add_paragraph()
        heading.add_run(label).bold = True

    if stype == "bullets":
        for item in content:
            doc.add_paragraph(item, style="List Bullet")
    elif stype == "activities":
        _render_activities(doc, content)
    elif stype == "exercises":
        for ex in content:
            p = doc.add_paragraph()
            p.add_run(ex.get("title", "Bài")).bold = True
            doc.add_paragraph(ex.get("content", ""))
    else:  # plain text — may be a list of lines or one string
        if isinstance(content, list):
            for line in content:
                doc.add_paragraph(str(line))
        else:
            doc.add_paragraph(str(content))


def _render_activities(doc: Document, activities: list[dict]) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    for i, text in enumerate(("Hoạt động", "Nội dung", "Thời lượng")):
        header[i].paragraphs[0].add_run(text).bold = True
    for act in activities:
        row = table.add_row().cells
        row[0].paragraphs[0].add_run(str(act.get("name", "")))
        row[1].paragraphs[0].add_run(f"Giáo viên: {act.get('teacher', '')}")
        row[1].add_paragraph().add_run(f"Học sinh: {act.get('students', '')}")
        row[2].paragraphs[0].add_run(f"{act.get('minutes', '')} phút")
