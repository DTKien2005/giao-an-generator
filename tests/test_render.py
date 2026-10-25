"""Renderers: Vietnamese .docx and .pdf from a structured document."""
from pathlib import Path

from app.config import GENERATED_DIR
from app.formats import default_format_spec, extract_format_spec, resolve_title
from app.render_docx import render_docx
from app.render_pdf import render_pdf

DOC = {
    "doc_type": "lesson_plan", "lesson": "Các số có hai chữ số", "week": "",
    "grade": "grade_1", "subject": "toan",
    "sections": {
        "objectives": ["Đọc, viết được các số từ 20 đến 50.",
                       "So sánh được các số có hai chữ số."],
        "materials": ["Giáo viên: que tính."],
        "activities": [
            {"name": "Khởi động", "minutes": 5,
             "teacher": "Tổ chức trò chơi đếm tiếp.",
             "students": "Học sinh đếm tiếp."},
            {"name": "Luyện tập", "minutes": 10,
             "teacher": "Giao bài tập.", "students": "Làm bảng con."},
        ],
        "assessment": "Nhận xét tinh thần học tập.",
        "appendix": "Đáp án: 38 < 41.",
    },
}


def test_docx_renders_vietnamese_and_tables(tmp_path: Path):
    spec = default_format_spec("lesson_plan")
    out = render_docx(DOC, spec, tmp_path / "t.docx")
    assert out.exists()
    import docx
    d = docx.Document(str(out))
    all_text = "\n".join(p.text for p in d.paragraphs)
    assert "GIÁO ÁN Toán LỚP 1" in all_text
    assert "So sánh được các số có hai chữ số." in all_text
    assert d.tables, "activities must render as a table"
    assert "Khởi động" in "\n".join(c.text for row in d.tables[0].rows for c in row.cells)


def test_pdf_renders_vietnamese(tmp_path: Path):
    spec = default_format_spec("lesson_plan")
    docx_path = render_docx(DOC, spec, tmp_path / "t.docx")
    out = render_pdf(DOC, spec, docx_path, tmp_path / "t.pdf")
    assert out.exists() and out.stat().st_size > 1000
    import pymupdf
    with pymupdf.open(str(out)) as pdf:
        text = pdf[0].get_text()
    assert "GIÁO ÁN Toán LỚP 1" in text
    assert "Đọc, viết được các số từ 20 đến 50." in text


def test_extract_format_spec_from_user_upload():
    sample = """GIÁO ÁN TOÁN LỚP 2
I. MỤC TIÊU
- Biết cộng hai số.
II. ĐỒ DÙNG DẠY HỌC
- Que tính.
III. CÁC HOẠT ĐỘNG DẠY HỌC
Hoạt động 1. Khởi động
"""
    spec = extract_format_spec(sample)
    assert spec is not None
    ids = [s["id"] for s in spec["sections"]]
    assert ids == ["objectives", "materials", "activities"]
    assert spec["source"] == "user_upload"


def test_extract_format_spec_returns_none_for_plain_text():
    assert extract_format_spec("chỉ là một đoạn văn bản thường không có mục gì") is None


def test_resolve_title_placeholders():
    title = resolve_title(default_format_spec("lesson_plan"), DOC)
    assert title == "GIÁO ÁN Toán LỚP 1"
    ws = resolve_title(default_format_spec("worksheet"),
                       {"grade": "grade_2", "subject": "toan", "week": "3"})
    assert ws == "PHIẾU BÀI TẬP TUẦN 3 – Toán LỚP 2"
